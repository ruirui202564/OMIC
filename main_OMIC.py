import os
import numpy as np
import openpyxl

from utils import utils_outputParam, data_load, config_params_OMIC
from run_OMIC import run_OMIC
from utils.utils import random_perm_n_old, random_perm_n
from utils.config_params_OMIC import create_param_OMIC as create_param_combination

path = os.path.realpath(__file__) 
result_addr = os.path.dirname(path)

#--------------All Variables--------------#
if __name__ == '__main__':
    __file__ = os.path.abspath('')

    method_name = "OMIC"
    data_type = 'syn' # 'binary', 'multi', 'syn'
    ratio = 0.50 # feature removal ratio

    data_name_list = []
    if data_type == "binary":
        data_name_list = ['internet_ads', 'SPAM', 'WEATHER', 'magic04', 'covtype', 'nomao',
                          'adult', 'Elec', 'PAKDD', 'creditcard', 'census', 'hepatitis']
    elif data_type == "multi":
        data_name_list = ['thyroid', 'GAS', 'shuttle', 'connect-4',
                          'OLYMPIC', 'covtype', 'lymph', 'zoo']
    elif data_type == "syn":
        data_name_list = ['static_RandomRBF', 'static_RandomTree',
                          'dynamic_RandomRBF','dynamic_RandomTree']

    result_all = np.zeros((len(data_name_list), 14))
    for data_index in range(len(data_name_list)):
        data_name = data_name_list[data_index]
        print("Data Name:", data_name)

        if data_type == 'binary':
            if data_name in ['covtype', 'magic04']:
                X, Y, X_haphazard, mask = data_load.data_load_imbalance_binary_other(data_name, ratio)
            else:
                X, Y, X_haphazard, mask = data_load.data_load_imbalance_binary(data_name, ratio)
        elif data_type == 'multi':
            X, Y, X_haphazard, mask = data_load.data_load_imbalance_multi(data_name, ratio)
        elif data_type == 'syn':
            X, Y, X_haphazard, mask = data_load.data_load_imbalance_multi_syn(data_name, ratio)

        param_dict = config_params_OMIC.config_OMIC()

        params_list = create_param_combination(param_dict)
        num_params = len(params_list[0])
        param_keys = list(params_list[0].keys())
        num_performance_metrics = 14
        result_all_params = np.zeros((len(params_list), num_performance_metrics + num_params + 1))
        for param_idx, params in enumerate(params_list):
            print(f"\n=== Processing parameter combination {param_idx + 1}/{len(params_list)} ===")
            print(f"Parameters: {params}")

            base_dir = os.path.join(result_addr, 'Results', f"{data_type}_{ratio:.2f}")
            result_addr_excel = os.path.join(base_dir, f"{data_name}.xlsx")
            result_merge_dataset = os.path.join(base_dir, "0 result_merge_dataset.xlsx")
            result_merge_param = os.path.join(base_dir, "0 result_merge_param.xlsx")
            result_merge_metric = os.path.join(base_dir, "0 result_merge_metric.xlsx")

            for file_path in [result_addr_excel, result_merge_dataset, result_merge_param, result_merge_metric]:
                dir_path = os.path.dirname(file_path)
                if not os.path.exists(dir_path):
                    os.makedirs(dir_path)

            accuracy = []
            bAccuracy = []
            precision = []
            f1 = []
            mGmean = []
            kappa = []
            time = []
            n_runs = 1

            n = X.shape[0]
            if data_type == 'syn':
                permutations = random_perm_n(n) # keep the original stream order to keep the imbalance ratio
            else:
                permutations = random_perm_n_old(n)
            for j in [0]: # range(10)
                print('permutation：', j + 1)
                perm = permutations[j]
                X_perm = X[perm]
                Y_perm = Y[perm]
                X_haphazard_perm = X_haphazard[perm]
                mask_perm = mask[perm]

                eval_list, dict_pred = run_OMIC(X_perm, Y_perm, X_haphazard_perm,
                                                mask_perm, n_runs, params)

                single_result = eval_list[0]
                accuracy.append(single_result['Accuracy'])
                bAccuracy.append(single_result['bAccuracy'])
                precision.append(single_result['Precision'])
                f1.append(single_result['F1'])
                mGmean.append(single_result['mGmean'])
                kappa.append(single_result['Kappa'])
                time.append(single_result['Time'])

                header = ['Accuracy','bAccuracy','Precision','F1','mGmean', 'Kappa', 'Time']
                sheet1 = "_".join(str(val) for val in list(params.values()))
                if os.path.exists(result_addr_excel):
                    workbook = openpyxl.load_workbook(result_addr_excel)
                    if sheet1 not in workbook.sheetnames:
                        sheet_one = workbook.create_sheet(sheet1)
                        for col_idx, header_name in enumerate(header, 1):
                            sheet_one.cell(row=1, column=col_idx).value = header_name
                    else:
                        sheet_one = workbook[sheet1]
                else:
                    workbook = openpyxl.Workbook()
                    sheet_one = workbook.active
                    sheet_one.title = sheet1
                    for col_idx, header_name in enumerate(header, 1):
                        sheet_one.cell(row=1, column=col_idx).value = header_name

                result_one = np.array([accuracy[j], bAccuracy[j], precision[j], f1[j], mGmean[j], kappa[j],
                                       time[j]])
                for col_idx, value in enumerate(result_one, 1):
                    sheet_one.cell(row=j + 2, column=col_idx).value = value

                workbook.save(result_addr_excel)

            results_dict, performance_results = utils_outputParam.calculate_mean_std_results(
                accuracy, bAccuracy, precision, f1, mGmean, kappa, time
            )
            result_all_params[param_idx, :num_performance_metrics] = performance_results
            for i, param_key in enumerate(param_keys):
                result_all_params[param_idx, num_performance_metrics + i] = params[param_key]
            result_all_params[param_idx, -1] = 1

            print(f'BAcc：{results_dict["bAccuracy_mean"]} F1：{results_dict["f1_mean"]}')

            utils_outputParam.save_mean_metrics_to_sheets(result_merge_metric, data_name, data_index,
                                        param_idx, params, results_dict)

            performance_headers = ['acc_m', 'acc_s', 'bAcc_m', 'bAcc_s', 'prescision_m', 'precision_s',
                                   'f1_m', 'f1_s', 'mGmean_m', 'mGmean_s', 'kappa_m', 'kappa_s', 'time_m', 'time_s']
            headers = ['dataset'] + param_keys + performance_headers
            sheet_merge_dataset = data_name
            if os.path.exists(result_merge_dataset):
                workbook = openpyxl.load_workbook(result_merge_dataset)
                if sheet_merge_dataset not in workbook.sheetnames:
                    sheet_merge = workbook.create_sheet(sheet_merge_dataset)
                    for col_idx, header_name in enumerate(headers, 1):
                        sheet_merge.cell(row=1, column=col_idx).value = header_name
                else:
                    sheet_merge = workbook[sheet_merge_dataset]
            else:
                workbook = openpyxl.Workbook()
                sheet_merge = workbook.active
                sheet_merge.title = sheet_merge_dataset
                for col_idx, header_name in enumerate(headers, 1):
                    sheet_merge.cell(row=1, column=col_idx).value = header_name

            row_num = param_idx + 2
            sheet_merge.cell(row=row_num, column=1).value = data_name
            for i, param_key in enumerate(param_keys):
                sheet_merge.cell(row=row_num, column=2 + i).value = params[param_key]
            start_col = 1 + num_params + 1
            for col_idx, value in enumerate(performance_results, start_col):
                sheet_merge.cell(row=row_num, column=col_idx).value = value
            workbook.save(result_merge_dataset)

            sheet_merge_param = "_".join(str(val) for val in list(params.values()))
            if os.path.exists(result_merge_param):
                workbook = openpyxl.load_workbook(result_merge_param)
                if sheet_merge_param not in workbook.sheetnames:
                    sheet_merge2 = workbook.create_sheet(sheet_merge_param)
                    for col_idx, header_name in enumerate(headers, 1):
                        sheet_merge2.cell(row=1, column=col_idx).value = header_name
                else:
                    sheet_merge2 = workbook[sheet_merge_param]
            else:
                workbook = openpyxl.Workbook()
                sheet_merge2 = workbook.active
                sheet_merge2.title = sheet_merge_param
                for col_idx, header_name in enumerate(headers, 1):
                    sheet_merge2.cell(row=1, column=col_idx).value = header_name

            row_num2 = data_index + 2
            sheet_merge2.cell(row=row_num2, column=1).value = data_name
            for i, param_key in enumerate(param_keys):
                sheet_merge2.cell(row=row_num2, column=2 + i).value = params[param_key]
            for col_idx, value in enumerate(performance_results, start_col):
                sheet_merge2.cell(row=row_num2, column=col_idx).value = value
            workbook.save(result_merge_param)