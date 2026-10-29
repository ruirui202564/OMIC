import numpy as np
import openpyxl
import os


def calculate_mean_std_results(accuracy, bAccuracy, precision, f1, mGmean, kappa, time):
    accuracy_mean = np.mean(accuracy)
    accuracy_std = np.std(accuracy)
    bAccuracy_mean = np.mean(bAccuracy)
    bAccuracy_std = np.std(bAccuracy)
    precision_mean = np.mean(precision)
    precision_std = np.std(precision)
    f1_mean = np.mean(f1)
    f1_std = np.std(f1)
    mGmean_mean = np.mean(mGmean)
    mGmean_std = np.std(mGmean)
    kappa_mean = np.mean(kappa)
    kappa_std = np.std(kappa)
    time_mean = np.mean(time)
    time_std = np.std(time)

    results_dict = {
        'accuracy_mean': accuracy_mean,
        'accuracy_std': accuracy_std,
        'bAccuracy_mean': bAccuracy_mean,
        'bAccuracy_std': bAccuracy_std,
        'precision_mean': precision_mean,
        'precision_std': precision_std,
        'f1_mean': f1_mean,
        'f1_std': f1_std,
        'mGmean_mean': mGmean_mean,
        'mGmean_std': mGmean_std,
        'kappa_mean': kappa_mean,
        'kappa_std': kappa_std,
        'time_mean': time_mean,
        'time_std': time_std
    }

    performance_results = np.array([
        accuracy_mean, accuracy_std, bAccuracy_mean, bAccuracy_std,
        precision_mean, precision_std, f1_mean, f1_std,
        mGmean_mean, mGmean_std, kappa_mean, kappa_std,
        time_mean, time_std
    ])

    return results_dict, performance_results


def save_mean_metrics_to_sheets(result_merge_metric, data_name, data_index,
                                param_idx, params, results_dict):
    dir_path = os.path.dirname(result_merge_metric)
    if not os.path.exists(dir_path):
        os.makedirs(dir_path)

    metric_mapping = {
        'acc_m': 'accuracy_mean',
        'bAcc_m': 'bAccuracy_mean',
        'precision_m': 'precision_mean',
        'f1_m': 'f1_mean',
        'mGmean_m': 'mGmean_mean',
        'kappa_m': 'kappa_mean',
        'time_m': 'time_mean'
    }

    param_combination_name = "_".join(str(val) for val in params.values())

    if os.path.exists(result_merge_metric):
        workbook = openpyxl.load_workbook(result_merge_metric)
    else:
        workbook = openpyxl.Workbook()
        if 'Sheet' in workbook.sheetnames:
            del workbook['Sheet']

    for sheet_name, metric_key in metric_mapping.items():
        if sheet_name not in workbook.sheetnames:
            sheet = workbook.create_sheet(sheet_name)
            sheet.cell(row=1, column=1).value = 'dataset'
            is_new_sheet = True
        else:
            sheet = workbook[sheet_name]
            is_new_sheet = False

        col_num = param_idx + 2

        if data_index == 0:
            sheet.cell(row=1, column=col_num).value = param_combination_name

        row_num = data_index + 2

        if is_new_sheet or sheet.cell(row=row_num, column=1).value is None:
            sheet.cell(row=row_num, column=1).value = data_name

        metric_value = results_dict[metric_key]
        sheet.cell(row=row_num, column=col_num).value = metric_value

    workbook.save(result_merge_metric)
    print(f"average results saved to: {result_merge_metric}")